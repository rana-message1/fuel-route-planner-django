"""
Import the supplied OPIS fuel-price file (.xlsx or .csv) into ``FuelStation``.

Pipeline
--------
1. Read rows and map the file's real headers (``OPIS Truckstop ID``,
   ``Truckstop Name``, ``Address``, ``City``, ``State``, ``Rack ID``,
   ``Retail Price``) to model fields.
2. Drop non-US rows (the file contains Canadian provinces: AB, ON, BC, ...).
3. Collapse duplicate OPIS IDs into one station, keeping the lowest price.
4. Geocode by city + state. The addresses are highway exits
   ("I-44, EXIT 283 & US-69"), so city-level coordinates are the practical
   precision. Sources, in priority order:
     a. ``data/station_geocode_overrides.csv`` (hand-curated / one-off Nominatim results)
     b. Census place centroid (``Place`` table)
     c. Census county-subdivision centroid (``Place`` table)
   ``--geocode-missing`` queries Nominatim for anything still unmatched and
   appends the results to the overrides file, so this happens only once.
5. Replace the table atomically.

This runs once at setup; API requests never touch the spreadsheet.
"""

from __future__ import annotations

import csv
import time
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from routing.models import FuelStation, Place
from routing.utils.places import US_STATES, normalize_place_name

COLUMN_MAP = {
    "opis truckstop id": "opis_id",
    "truckstop name": "name",
    "address": "address",
    "city": "city",
    "state": "state",
    "rack id": "rack_id",
    "retail price": "retail_price",
}
REQUIRED_FIELDS = {"opis_id", "name", "address", "city", "state", "retail_price"}
PRICE_PRECISION = Decimal("0.0001")
OVERRIDE_COLUMNS = ["state", "city", "latitude", "longitude", "source"]


@dataclass
class StationRow:
    opis_id: int
    name: str
    address: str
    city: str
    state: str
    rack_id: int | None
    retail_price: Decimal
    observations: int = 1


def _normalize_header(value) -> str:
    return " ".join(str(value or "").split()).lower()


def read_rows(path: Path) -> list[dict[str, object]]:
    """Return rows as dicts keyed by model field name."""
    if not path.exists():
        raise CommandError(f"File not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            table = list(csv.reader(handle))
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            table = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        raise CommandError(f"Unsupported file type '{suffix}'. Use .xlsx or .csv.")

    if not table:
        raise CommandError(f"{path.name} is empty.")
    header = [COLUMN_MAP.get(_normalize_header(cell)) for cell in table[0]]
    missing = REQUIRED_FIELDS - set(header)
    if missing:
        raise CommandError(
            f"{path.name} is missing required columns for: {sorted(missing)}. Found headers: {table[0]}"
        )
    return [
        {field: value for field, value in zip(header, row, strict=False) if field}
        for row in table[1:]
        if any(cell not in (None, "") for cell in row)
    ]


def _to_int(value) -> int:
    return int(Decimal(str(value).strip()))


def parse_row(raw: dict[str, object]) -> StationRow:
    """Validate one row; raises ``ValueError`` for unusable data."""
    try:
        opis_id = _to_int(raw["opis_id"])
        price = Decimal(str(raw["retail_price"]).strip()).quantize(PRICE_PRECISION)
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(
            f"invalid id or price ({raw.get('opis_id')!r}, {raw.get('retail_price')!r})"
        ) from exc
    if price <= 0:
        raise ValueError(f"non-positive price {price}")

    rack_raw = raw.get("rack_id")
    try:
        rack_id = _to_int(rack_raw) if rack_raw not in (None, "") else None
    except (InvalidOperation, ValueError):
        rack_id = None

    return StationRow(
        opis_id=opis_id,
        name=" ".join(str(raw["name"] or "").split()),
        address=" ".join(str(raw["address"] or "").split()),
        city=" ".join(str(raw["city"] or "").split()),
        state=str(raw["state"] or "").strip().upper(),
        rack_id=rack_id,
        retail_price=price,
    )


def load_overrides(path: Path) -> dict[tuple[str, str], tuple[float, float]]:
    if not path.exists():
        return {}
    overrides = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            key = (row["state"].strip().upper(), normalize_place_name(row["city"]))
            overrides[key] = (float(row["latitude"]), float(row["longitude"]))
    return overrides


class Command(BaseCommand):
    help = "Import the fuel-price spreadsheet (.xlsx or .csv) into the FuelStation table."

    def add_arguments(self, parser):
        config = settings.FUEL_PLANNER
        parser.add_argument("file", nargs="?", type=Path, default=config["FUEL_PRICES_FILE"])
        parser.add_argument("--overrides-file", type=Path, default=config["STATION_GEOCODE_OVERRIDES_FILE"])
        parser.add_argument(
            "--geocode-missing",
            action="store_true",
            help="Geocode unmatched cities with Nominatim (1 req/s) and save them to the overrides file.",
        )

    def handle(self, *args, **options):
        started = time.perf_counter()
        path = Path(options["file"])
        if not Place.objects.exists():
            raise CommandError("The Place table is empty. Run `python manage.py import_places` first.")

        raw_rows = read_rows(path)
        stations, stats = self._aggregate(raw_rows)
        overrides_path = Path(options["overrides_file"])
        sources = self._geocode(stations, overrides_path, options["geocode_missing"])

        with transaction.atomic():
            FuelStation.objects.all().delete()
            FuelStation.objects.bulk_create(
                [
                    FuelStation(
                        opis_id=row.opis_id,
                        name=row.name,
                        address=row.address,
                        city=row.city,
                        state=row.state,
                        rack_id=row.rack_id,
                        retail_price=row.retail_price,
                        price_observations=row.observations,
                        latitude=coords[0] if coords else None,
                        longitude=coords[1] if coords else None,
                        geocode_source=source,
                    )
                    for row, (coords, source) in zip(stations, sources, strict=True)
                ],
                batch_size=1000,
            )

        by_source = Counter(source for _, source in sources)
        elapsed = time.perf_counter() - started
        self.stdout.write(f"Rows read:                    {len(raw_rows):,}")
        self.stdout.write(f"Rows skipped (invalid):       {stats['invalid']:,}")
        self.stdout.write(f"Rows skipped (outside USA):   {stats['non_us']:,}")
        self.stdout.write(f"Duplicate rows merged:        {stats['duplicates']:,}")
        self.stdout.write(f"Stations imported:            {len(stations):,}")
        for source, count in sorted(by_source.items()):
            self.stdout.write(f"  geocoded via {source:<27} {count:,}")
        if by_source.get(FuelStation.GeocodeSource.NONE):
            self.stdout.write(
                self.style.WARNING(
                    "Some stations have no coordinates and will never be suggested. "
                    "Run with --geocode-missing to resolve them."
                )
            )
        self.stdout.write(self.style.SUCCESS(f"Done in {elapsed:.1f}s."))

    def _aggregate(self, raw_rows):
        stats = Counter()
        by_id: dict[int, StationRow] = {}
        for line_number, raw in enumerate(raw_rows, start=2):
            try:
                row = parse_row(raw)
            except ValueError as exc:
                stats["invalid"] += 1
                self.stderr.write(f"Skipping line {line_number}: {exc}")
                continue
            if row.state not in US_STATES:
                stats["non_us"] += 1
                continue
            existing = by_id.get(row.opis_id)
            if existing is None:
                by_id[row.opis_id] = row
                continue
            # Same OPIS ID listed again: keep the cheapest listing's details.
            stats["duplicates"] += 1
            observations = existing.observations + 1
            if row.retail_price < existing.retail_price:
                by_id[row.opis_id] = row
            by_id[row.opis_id].observations = observations
        return list(by_id.values()), stats

    def _geocode(self, stations: list[StationRow], overrides_path: Path, geocode_missing: bool):
        overrides = load_overrides(overrides_path)
        states = {row.state for row in stations}
        places = {
            (state, normalized): (lat, lon, kind)
            for state, normalized, lat, lon, kind in Place.objects.filter(state__in=states).values_list(
                "state", "normalized_name", "latitude", "longitude", "kind"
            )
        }
        kind_to_source = {
            Place.Kind.PLACE: FuelStation.GeocodeSource.CENSUS_PLACE,
            Place.Kind.COUNTY_SUBDIVISION: FuelStation.GeocodeSource.CENSUS_COUNTY_SUBDIVISION,
        }

        def lookup(row: StationRow):
            key = (row.state, normalize_place_name(row.city))
            if key in overrides:
                return overrides[key], FuelStation.GeocodeSource.OVERRIDE
            if key in places:
                lat, lon, kind = places[key]
                return (lat, lon), kind_to_source[kind]
            return None, FuelStation.GeocodeSource.NONE

        results = [lookup(row) for row in stations]
        if geocode_missing:
            unmatched = sorted(
                {
                    (row.state, row.city)
                    for row, (coords, _) in zip(stations, results, strict=True)
                    if coords is None
                }
            )
            if unmatched and self._geocode_with_nominatim(unmatched, overrides_path):
                overrides = load_overrides(overrides_path)
                results = [lookup(row) for row in stations]
        return results

    def _geocode_with_nominatim(self, cities: list[tuple[str, str]], overrides_path: Path) -> int:
        from routing.exceptions import GeocodingUnavailableError
        from routing.services.geocoding_service import NominatimClient

        client = NominatimClient()
        self.stdout.write(f"Geocoding {len(cities)} unmatched cities with Nominatim (1 request/second)...")
        new_rows = []
        for state, city in cities:
            try:
                result = client.search(city=city, state=US_STATES[state])
            except GeocodingUnavailableError as exc:
                self.stderr.write(f"  {city}, {state}: {exc.detail}")
                result = None
            if result:
                new_rows.append(
                    {
                        "state": state,
                        "city": city,
                        "latitude": f"{result[0]:.6f}",
                        "longitude": f"{result[1]:.6f}",
                        "source": "nominatim",
                    }
                )
            else:
                self.stderr.write(f"  {city}, {state}: not found")
            time.sleep(1.1)  # Nominatim usage policy: max 1 request per second

        if new_rows:
            write_header = not overrides_path.exists()
            overrides_path.parent.mkdir(parents=True, exist_ok=True)
            with overrides_path.open("a", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=OVERRIDE_COLUMNS)
                if write_header:
                    writer.writeheader()
                writer.writerows(new_rows)
            self.stdout.write(f"Saved {len(new_rows)} coordinates to {overrides_path}.")
        return len(new_rows)
